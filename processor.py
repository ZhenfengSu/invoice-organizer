"""Extract, validate and package invoice PDFs."""

from __future__ import annotations

import hashlib
import json
import os
import re
import shutil
import unicodedata
import urllib.error
import urllib.request
import zipfile
from collections import Counter, defaultdict
from copy import copy
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path, PurePosixPath
from typing import Iterable

import pymupdf as fitz
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


CATEGORIES = ("机票", "高铁票", "餐饮", "打车", "住宿")
MAX_PAGES = 10
MONEY = r"(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d{1,2})?"
USCC_ALPHABET = "0123456789ABCDEFGHJKLMNPQRTUWXY"
USCC_WEIGHTS = (1, 3, 9, 27, 19, 26, 16, 17, 20, 29, 25, 13, 8, 24, 10, 30, 28)
OCR_ENGINE = None


@dataclass
class Invoice:
    path: Path
    source_name: str
    region: str = "未识别"
    category: str = ""
    invoice_number: str = ""
    buyer_name: str = ""
    tax_id: str = ""
    amount: Decimal | None = None
    currency: str = ""
    issues: list[str] = field(default_factory=list)
    method: str = ""


def normalize_name(value: str) -> str:
    return re.sub(r"\s", "", unicodedata.normalize("NFKC", value or "")).casefold()


def normalize_tax(value: str) -> str:
    return re.sub(r"\s", "", value or "").upper()


def valid_uscc(value: str) -> bool:
    value = normalize_tax(value)
    if len(value) != 18 or any(c not in USCC_ALPHABET for c in value):
        return False
    check = (31 - sum(USCC_ALPHABET.index(c) * weight for c, weight in zip(value[:17], USCC_WEIGHTS)) % 31) % 31
    return value[-1] == USCC_ALPHABET[check]


def _ocr_page(page: fitz.Page) -> tuple[str, str]:
    global OCR_ENGINE
    if OCR_ENGINE is None:
        from rapidocr_onnxruntime import RapidOCR

        OCR_ENGINE = RapidOCR()
    pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
    result, _ = OCR_ENGINE(pix.tobytes("png"))
    if not result:
        return "", ""
    items = []
    for box, text, _score in result:
        x = min(point[0] for point in box)
        y = min(point[1] for point in box)
        items.append((y, x, text))
    items.sort(key=lambda item: (round(item[0] / 18), item[1]))
    all_text = "\n".join(item[2] for item in items)
    left_text = "\n".join(item[2] for item in items if item[1] < pix.width * 0.54)
    return all_text, left_text


def read_pdf(path: Path) -> tuple[str, str, str]:
    full, left, methods = [], [], []
    with fitz.open(path) as document:
        if not document or len(document) > MAX_PAGES:
            raise ValueError(f"PDF 页数须在 1 至 {MAX_PAGES} 页之间")
        for page in document:
            text = page.get_text("text", sort=True)
            side = page.get_text("text", sort=True, clip=fitz.Rect(0, 0, page.rect.width * 0.54, page.rect.height))
            if len(re.sub(r"\s", "", text)) < 40:
                text, side = _ocr_page(page)
                methods.append("OCR")
            else:
                methods.append("PDF 文字")
            full.append(text)
            left.append(side)
    return "\n".join(full), "\n".join(left), "+".join(sorted(set(methods)))


def _first(pattern: str, text: str, flags: int = 0) -> str:
    match = re.search(pattern, text, flags)
    return match.group(1).strip() if match else ""


def _money(value: str) -> Decimal | None:
    try:
        amount = Decimal(value.replace(",", ""))
        return amount.quantize(Decimal("0.01")) if amount >= 0 else None
    except InvalidOperation:
        return None


def domestic_amount(text: str) -> Decimal | None:
    compact = re.sub(r"[ \t]", "", text)
    patterns = [
        rf"价税合计[^\n]{{0,45}}?(?:小写|\(小写\)|（小写）)[^\d\n]{{0,8}}({MONEY})",
        rf"(?:小写|\(小写\)|（小写）)[^\d\n]{{0,8}}({MONEY})",
        rf"(?:票价|退票费)[：:][^\d\n]{{0,8}}({MONEY})",
    ]
    for pattern in patterns:
        found = _first(pattern, compact, re.I)
        if found:
            return _money(found)
    return None


def overseas_amount(text: str) -> tuple[Decimal | None, str]:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    labels = r"(?:grand\s*total|total\s*(?:paid|due|amount)?|amount\s*(?:paid|due)|balance\s*paid|payment\s*total|合计|總計|合計金額)"
    excluded = r"sub\s*total|subtotal|tax|tip|change|discount|saving"
    for line in reversed(lines):
        if re.search(labels, line, re.I) and not re.search(excluded, line, re.I):
            candidates = list(re.finditer(rf"(?<!\d)({MONEY})(?!\d)", line))
            if candidates:
                amount = _money(candidates[-1].group(1))
                if amount is not None:
                    line_currency = detect_currency(line)
                    return amount, line_currency if line_currency != "未标明" else detect_currency(text)
    # A lone currency amount is useful for minimal receipts, but multiple amounts are ambiguous.
    candidates = re.findall(rf"(?:USD|SGD|EUR|GBP|JPY|HKD|AUD|CAD|CNY|US\$|S\$|HK\$|A\$|C\$|[$€£¥￥])\s*({MONEY})", text, re.I)
    if len(candidates) == 1:
        return _money(candidates[0]), detect_currency(text)
    return None, detect_currency(text)


def detect_currency(text: str) -> str:
    for pattern, currency in (
        (r"\bSGD\b|(?<![A-Za-z])S\$", "SGD"), (r"\bUSD\b|US\$", "USD"),
        (r"\bHKD\b|HK\$", "HKD"), (r"\bAUD\b|(?<![A-Za-z])A\$", "AUD"),
        (r"\bCAD\b|(?<![A-Za-z])C\$", "CAD"), (r"\bEUR\b|€", "EUR"),
        (r"\bGBP\b|£", "GBP"), (r"\bJPY\b", "JPY"),
        (r"\bCNY\b|￥", "CNY"),
    ):
        if re.search(pattern, text, re.I):
            return currency
    return "未标明"


def classify(text: str) -> tuple[str, str]:
    condensed = re.sub(r"\s", "", text)
    domestic = bool(re.search(r"电子发票|发票号码|国家税务总局|税务局|统一社会信用代码|纳税人识别号", condensed))
    if domestic:
        if re.search(r"铁路电子客票|电子客票号|高铁|火车票|\b[GDCKTZ]\d{2,5}\b", text, re.I):
            return "国内", "高铁票"
        if re.search(r"民航|航空|机票|航班|客票款|燃油附加费|行程单", condensed) or ("旅行社" in condensed and "退票费" in condensed):
            return "国内", "机票"
        if re.search(r"住宿|酒店|宾馆|旅馆|客房|民宿", condensed):
            return "国内", "住宿"
        if re.search(r"餐饮|餐费|饭店|食府|酒楼|餐厅", condensed):
            return "国内", "餐饮"
        if re.search(r"出租车|网约车|打车|约车|客运服务费|旅客运输服务", condensed):
            return "国内", "打车"
        return "国内", ""
    if re.search(r"\breceipt\b|\binvoice\b|\bgrand\s*total\b|\btotal\b|\bamount\s*paid\b", text, re.I):
        return "海外", ""
    return "未识别", ""


def _llm_classify(text: str) -> tuple[str, str]:
    key = os.getenv("LLM_API_KEY", "")
    base_url = os.getenv("LLM_BASE_URL", "")
    if not key or not base_url:
        return "", ""
    prompt = (
        "只判断以下票据的地域和费用类型。返回单个 JSON 对象，字段 region 只能是 国内、海外、未识别，"
        "category 只能是 机票、高铁票、餐饮、打车、住宿、空字符串。不要猜测不在文本中的信息。\n"
        + text[:8000]
    )
    payload = json.dumps({"model": os.getenv("LLM_MODEL", "gpt-5.6-luna"), "input": prompt, "max_output_tokens": 150}).encode()
    request = urllib.request.Request(
        base_url.rstrip("/") + "/responses",
        data=payload,
        headers={"Authorization": "Bearer " + key, "Content-Type": "application/json"},
    )
    try:
        with urllib.request.urlopen(request, timeout=30) as response:
            data = json.load(response)
        output = data.get("output_text") or " ".join(
            part.get("text", "") for item in data.get("output", [])
            for part in item.get("content", []) if part.get("type") == "output_text"
        )
        parsed = json.loads(re.search(r"\{.*?\}", output, re.S).group())
        region = parsed.get("region", "")
        category = parsed.get("category", "")
        if region in ("国内", "海外", "未识别") and category in (*CATEGORIES, ""):
            return region, category
    except (urllib.error.URLError, ValueError, KeyError, AttributeError):
        pass
    return "", ""


def extract(path: Path, source_name: str, buyer_names: list[str], tax_ids: list[str], use_llm: bool) -> Invoice:
    item = Invoice(path=path, source_name=source_name)
    try:
        text, left, item.method = read_pdf(path)
    except Exception as exc:
        item.issues.append(f"无法读取 PDF：{exc}")
        return item
    if not text.strip():
        item.issues.append("PDF 未识别出文字")
        return item
    item.region, item.category = classify(text)
    if use_llm and (item.region == "未识别" or item.region == "国内" and not item.category):
        region, category = _llm_classify(text)
        if region:
            item.region, item.category = region, category
    if item.region == "国内":
        item.invoice_number = _first(r"发票号码\s*[：:]?\s*(\d{8,20})", text)
        buyer_text = left or text
        item.buyer_name = _first(r"(?:购买方名称|购方名称|名称)\s*[：:]\s*([^\s：:]{2,60})", buyer_text)
        tax_pattern = r"(?:统一社会信用代码\s*/\s*纳税人识别号|统一社会信用代码|纳税人识别号)\s*[：:]\s*([0-9A-Z]{15,20})"
        # Some layouts put both tax labels on one line and their values on the next.
        label = re.search(r"(?:统一社会信用代码\s*/\s*纳税人识别号|统一社会信用代码|纳税人识别号)\s*[：:]", buyer_text)
        following = buyer_text[label.end():label.end() + 180] if label else ""
        next_line_code = _first(r"(?<![0-9A-Z])([0-9A-Z]{18}|[0-9A-Z]{15})(?![0-9A-Z])", following, re.I)
        item.tax_id = normalize_tax(
            _first(tax_pattern, buyer_text, re.I)
            or next_line_code
            or _first(tax_pattern, text, re.I)
        )
        item.amount = domestic_amount(text)
        item.currency = "CNY"
        if not item.category:
            item.issues.append("未能归入五类国内发票")
        if not item.invoice_number:
            item.issues.append("缺少有效发票号码")
        if item.amount is None:
            item.issues.append("未能识别价税合计金额")
        if not item.buyer_name:
            item.issues.append("未能识别购买方名称")
        elif buyer_names and normalize_name(item.buyer_name) not in {normalize_name(x) for x in buyer_names}:
            item.issues.append("购买方名称与设定值不一致")
        if not item.tax_id:
            item.issues.append("未能识别购买方税号")
        else:
            if len(item.tax_id) == 18 and not valid_uscc(item.tax_id):
                item.issues.append("统一社会信用代码校验位错误")
            if tax_ids and item.tax_id not in {normalize_tax(x) for x in tax_ids}:
                item.issues.append("购买方税号与设定值不一致")
        if buyer_names and len(buyer_names) == len(tax_ids) and item.buyer_name and item.tax_id:
            allowed_pairs = {(normalize_name(name), normalize_tax(tax)) for name, tax in zip(buyer_names, tax_ids)}
            if (normalize_name(item.buyer_name), item.tax_id) not in allowed_pairs:
                if not any("与设定值不一致" in issue for issue in item.issues):
                    item.issues.append("购买方名称与税号组合不匹配")
    elif item.region == "海外":
        item.amount, item.currency = overseas_amount(text)
        if item.amount is None:
            item.issues.append("未能可靠识别海外收据总金额")
        if item.currency == "未标明":
            item.issues.append("海外收据未标明币种")
    else:
        item.issues.append("无法判断是中国发票还是海外收据")
    return item


def _workbook(headers: list[str], rows: list[list], title: str, totals: list[list] | None = None) -> Workbook:
    book = Workbook()
    sheet = book.active
    sheet.title = title
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    if totals:
        for row in totals:
            sheet.append(row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, len(rows) + 1)}"
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="163A42")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 28
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if isinstance(cell.value, (int, float)) and "金额" in headers[cell.column - 1]:
                cell.number_format = '#,##0.00'
            if cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F0F5F3")
    for column in sheet.columns:
        values = [len(str(cell.value or "")) for cell in column]
        sheet.column_dimensions[column[0].column_letter].width = min(max(max(values, default=8) * 1.6, 14), 65)
    return book


def _amount(value: Decimal | None) -> float | None:
    return float(value) if value is not None else None


def _safe_destination(folder: Path, filename: str) -> Path:
    stem = Path(filename).stem[:100] or "票据"
    suffix = Path(filename).suffix.lower() or ".pdf"
    candidate = folder / f"{stem}{suffix}"
    number = 2
    while candidate.exists():
        candidate = folder / f"{stem}-{number}{suffix}"
        number += 1
    return candidate


def package(items: list[Invoice], output_dir: Path) -> tuple[Path, dict]:
    output_dir.mkdir(parents=True, exist_ok=True)
    counts = Counter(item.invoice_number for item in items if item.region == "国内" and item.invoice_number)
    hashes = defaultdict(list)
    for item in items:
        hashes[hashlib.sha256(item.path.read_bytes()).hexdigest()].append(item)
    for item in items:
        if item.region == "国内" and item.invoice_number and counts[item.invoice_number] > 1:
            item.issues.append("重复发票号码（同号票据全部排除）")
        if len(hashes[hashlib.sha256(item.path.read_bytes()).hexdigest()]) > 1 and "重复文件内容（全部排除）" not in item.issues:
            item.issues.append("重复文件内容（全部排除）")

    domestic_dir = output_dir / "国内发票"
    overseas_dir = output_dir / "海外收据"
    error_dir = output_dir / "错误发票"
    for folder in (domestic_dir, overseas_dir, error_dir):
        folder.mkdir(exist_ok=True)
    domestic = {category: [] for category in CATEGORIES}
    overseas, errors = [], []
    for item in items:
        original_name = PurePosixPath(item.source_name.replace("\\", "/")).name
        if item.issues:
            shutil.copy2(item.path, _safe_destination(error_dir, original_name))
            errors.append(item)
        elif item.region == "国内":
            folder = domestic_dir / item.category
            folder.mkdir(exist_ok=True)
            shutil.copy2(item.path, _safe_destination(folder, original_name))
            domestic[item.category].append(item)
        else:
            shutil.copy2(item.path, _safe_destination(overseas_dir, original_name))
            overseas.append(item)

    domestic_book = _workbook(["类型", "张数", "金额合计 (CNY)"], [
        [category, len(domestic[category]), float(sum((x.amount for x in domestic[category]), Decimal("0")))]
        for category in CATEGORIES
    ], "总览", [["国内总计", sum(len(group) for group in domestic.values()), float(sum((x.amount for group in domestic.values() for x in group), Decimal("0")))]] )
    for category in CATEGORIES:
        rows = [[item.source_name, item.invoice_number, _amount(item.amount)] for item in domestic[category]]
        total = float(sum((item.amount for item in domestic[category]), Decimal("0")))
        sheet = domestic_book.create_sheet(category)
        styled = _workbook(["文件名", "发票号码", "金额 (CNY)"], rows, category, [["类型合计", "", total]]).active
        for row in styled:
            for cell in row:
                sheet[cell.coordinate] = cell.value
                if cell.has_style:
                    sheet[cell.coordinate]._style = copy(cell._style)
        for key, dim in styled.column_dimensions.items():
            sheet.column_dimensions[key].width = dim.width
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = styled.auto_filter.ref
        category_dir = domestic_dir / category
        category_dir.mkdir(exist_ok=True)
        _workbook(["文件名", "发票号码", "金额 (CNY)"], rows, category, [["类型合计", "", total]]).save(category_dir / f"{category}统计.xlsx")
    domestic_book.save(domestic_dir / "国内发票汇总.xlsx")

    overseas_totals = defaultdict(lambda: Decimal("0"))
    for item in overseas:
        overseas_totals[item.currency] += item.amount
    overseas_rows = [[item.source_name, _amount(item.amount), item.currency] for item in overseas]
    overseas_book = _workbook(["文件名", "金额", "币种"], overseas_rows, "海外收据", [
        [f"{currency} 总计", float(total), currency] for currency, total in sorted(overseas_totals.items())
    ])
    overseas_book.save(overseas_dir / "海外收据汇总.xlsx")

    error_rows = [[item.source_name, "；".join(item.issues), item.region, item.category,
                   item.invoice_number, item.buyer_name, item.tax_id, _amount(item.amount), item.currency]
                  for item in errors]
    _workbook(["文件名", "错误原因", "地区", "类别", "发票号码", "购买方名称", "购买方税号", "金额", "币种"],
              error_rows, "错误报告").save(error_dir / "错误报告.xlsx")
    (error_dir / "错误报告.txt").write_text(
        "错误发票报告\n\n" + "\n".join(f"{item.source_name}\n  {'；'.join(item.issues)}" for item in errors)
        + ("\n" if errors else "无错误发票。\n"), encoding="utf-8"
    )

    archive = output_dir.parent / "发票分类结果.zip"
    with zipfile.ZipFile(archive, "w", zipfile.ZIP_DEFLATED) as z:
        for file in sorted(output_dir.rglob("*")):
            if file.is_file():
                z.write(file, file.relative_to(output_dir))
    summary = {
        "processed": len(items), "domestic": sum(map(len, domestic.values())),
        "overseas": len(overseas), "errors": len(errors),
        "domestic_total": f"{sum((x.amount for group in domestic.values() for x in group), Decimal('0')):.2f}",
        "overseas_totals": {key: f"{value:.2f}" for key, value in sorted(overseas_totals.items())},
        "categories": {key: len(value) for key, value in domestic.items()},
        "error_details": [{"file": item.source_name, "reasons": item.issues} for item in errors],
    }
    return archive, summary


def process(files: Iterable[tuple[Path, str]], output_dir: Path, buyer_names: list[str], tax_ids: list[str], use_llm: bool = False) -> tuple[Path, dict]:
    records = [extract(path, name, buyer_names, tax_ids, use_llm) for path, name in files]
    return package(records, output_dir)
